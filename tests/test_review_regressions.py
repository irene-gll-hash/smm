import asyncio
import io
import unittest
import tempfile
from pathlib import Path
import zipfile
from unittest.mock import AsyncMock, patch

from docx import Document

from qifa_smm.db import Database
from qifa_smm.extractors import ContentExtractor
from qifa_smm.telegram import TelegramController


class OfficeExtractionTests(unittest.TestCase):
    def extract_archive(self, name, files):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            for path, content in files.items():
                archive.writestr(path, content)
        return ContentExtractor().from_bytes(name, 'application/octet-stream', buffer.getvalue())

    def test_xlsx_resolves_shared_strings_in_cell_order(self):
        text, readable = self.extract_archive('sample.xlsx', {
            'xl/sharedStrings.xml': '<sst><si><t>unused</t></si><si><t>Revenue</t></si></sst>',
            'xl/worksheets/sheet1.xml': '<worksheet><sheetData><row><c r="A1" t="s"><v>1</v></c><c r="C1"><v>120</v></c><c r="D1" t="inlineStr"><is><t>RUB</t></is></c></row></sheetData></worksheet>',
        })
        self.assertTrue(readable)
        self.assertIn('A1: Revenue\tC1: 120\tD1: RUB', text)
        self.assertNotIn('unused', text)

    def test_empty_spreadsheet_is_unreadable(self):
        self.assertEqual(self.extract_archive('empty.xlsx', {
            'xl/worksheets/sheet1.xml': '<worksheet><sheetData/></worksheet>',
        }), ('', False))

    def test_slides_follow_numeric_order_and_exclude_notes(self):
        text, readable = self.extract_archive('sample.pptx', {
            'ppt/slides/slide10.xml': '<slide><t>ten</t></slide>',
            'ppt/slides/slide2.xml': '<slide><t>two</t></slide>',
            'ppt/slides/slide1.xml': '<slide><t>one</t></slide>',
        })
        self.assertEqual(text, 'one\ntwo\nten')
        self.assertTrue(readable)

    def test_docx_table_retains_position_and_cell_values(self):
        doc = Document()
        doc.add_paragraph('Before')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Revenue'
        table.cell(0, 1).text = '120'
        doc.add_paragraph('After')
        buffer = io.BytesIO()
        doc.save(buffer)
        text, readable = ContentExtractor().from_bytes('table.docx', '', buffer.getvalue())
        self.assertEqual(text, 'Before\nRevenue\t120\nAfter')
        self.assertTrue(readable)


class CallbackTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / 'state.sqlite')
        await self.db.initialize()
        self.batch = [
            {'update_id': 1, 'callback_query': {'id': 'old', 'from': {'id': 1}, 'data': 'm:status'}},
            {'update_id': 2, 'message': {'from': {'id': 1}, 'chat': {'id': 1}, 'text': '/status'}},
        ]

    async def asyncTearDown(self):
        self.tmp.cleanup()

    async def test_failed_callback_ack_does_not_block_updates(self):
        api = AsyncMock()
        api.updates.side_effect = [self.batch, asyncio.CancelledError()]
        api.answer_callback.side_effect = RuntimeError('expired')
        controller = TelegramController(None, api, self.db, None)
        controller.handle = AsyncMock()
        with self.assertLogs('qifa_smm.telegram', level='WARNING'):
            with self.assertRaises(asyncio.CancelledError):
                await controller.run()
        self.assertEqual(controller.handle.await_count, 2)
        self.assertEqual(await self.db.pending_telegram_updates(), [])
        self.assertEqual(await self.db.get_setting('telegram_offset'), '3')
        self.assertTrue(await self.db.is_update_processed(1))

    async def test_handler_failure_survives_restart_and_batch_is_preserved(self):
        api = AsyncMock()
        api.updates.return_value = self.batch
        controller = TelegramController(None, api, self.db, None)
        controller.handle = AsyncMock(side_effect=RuntimeError('temporary outage'))
        with patch('qifa_smm.telegram.asyncio.sleep', AsyncMock(side_effect=asyncio.CancelledError())):
            with self.assertLogs('qifa_smm.telegram', level='ERROR'):
                with self.assertRaises(asyncio.CancelledError):
                    await controller.run()
        self.assertFalse(await self.db.is_update_processed(1))
        self.assertEqual(await self.db.pending_telegram_updates(), self.batch)
        restarted = Database(self.db.path)
        await restarted.initialize()
        api = AsyncMock()
        api.updates.side_effect = asyncio.CancelledError()
        controller = TelegramController(None, api, restarted, None)
        controller.handle = AsyncMock()
        with self.assertRaises(asyncio.CancelledError):
            await controller.run()
        self.assertEqual(controller.handle.await_count, 2)
        api.updates.assert_awaited_once_with(3)
        self.assertEqual(await restarted.pending_telegram_updates(), [])

    async def test_repeated_batch_does_not_repeat_processed_events(self):
        await self.db.save_telegram_updates(self.batch)
        await self.db.mark_update_processed(1)
        await self.db.save_telegram_updates(self.batch)
        self.assertEqual(await self.db.pending_telegram_updates(), self.batch[1:])

    async def test_batch_storage_and_offset_are_atomic(self):
        with self.assertRaises(KeyError):
            await self.db.save_telegram_updates([self.batch[0], {}])
        self.assertEqual(await self.db.pending_telegram_updates(), [])
        self.assertEqual(await self.db.get_setting('telegram_offset', '0'), '0')

    async def test_explicit_rejection_is_reported_and_does_not_block_queue(self):
        from qifa_smm.security import UserFacingError
        api = AsyncMock()
        api.updates.side_effect = [self.batch, asyncio.CancelledError()]
        controller = TelegramController(None, api, self.db, None)
        controller.handle = AsyncMock(side_effect=[UserFacingError('Версия устарела'), None])
        with self.assertRaises(asyncio.CancelledError):
            await controller.run()
        api.send_message.assert_awaited_once_with(1, 'Версия устарела')
        self.assertEqual(controller.handle.await_count, 2)
        self.assertEqual(await self.db.pending_telegram_updates(), [])

    async def test_failed_enqueue_keeps_period_selection_for_retry(self):
        api = AsyncMock()
        from qifa_smm.domain import TelegramEvent
        event = TelegramEvent(update_id=5, user_id=1, chat_id=1, kind='callback', callback_data='per:0')
        await self.db.set_session(1, 1, 'period_select', 'materials', {'choices': [{'period_id': 'P'}]})
        controller = TelegramController(None, api, self.db, None)
        with patch.object(self.db, 'enqueue_job', AsyncMock(side_effect=RuntimeError('busy'))):
            with self.assertRaises(RuntimeError):
                await controller._callback(event)
        self.assertIsNotNone(await self.db.get_session(1))
        await controller._callback(event)
        self.assertIsNone(await self.db.get_session(1))
        self.assertEqual((await self.db.claim_next_job()).payload, {'period_id': 'P'})
